"""Tests voor NOVA zonder Windows: nagebootste ECO-server en Ollama. Draaien: python -m unittest -v test_nova"""
import json
import sys
import tempfile
import threading
import types
import unittest
from unittest import mock
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import numpy as np

import nova_core as core
import nova_pc as pc
import nova_stem as stem

GEHEIM = "geheim"


class NepServer:
    """Kleine HTTP-server die ECO's functie of Ollama nabootst."""

    def __init__(self, behandel):
        nep = self

        class H(BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def antwoord(self, code, body, soort="application/json", headers=None):
                data = body if isinstance(body, bytes) else json.dumps(body).encode()
                self.send_response(code)
                self.send_header("content-type", soort)
                for k, v in (headers or {}).items():
                    self.send_header(k, v)
                self.send_header("content-length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def do_GET(self):
                behandel(self, "GET", None)

            def do_POST(self):
                n = int(self.headers.get("content-length") or 0)
                raw = self.rfile.read(n) if n else b""
                try:
                    body = json.loads(raw) if raw and "json" in (self.headers.get("content-type") or "") else raw
                except ValueError:
                    body = raw
                nep.verzoeken.append((self.path, self.headers.get("x-jarvis-key"), body))
                behandel(self, "POST", body)

        self.verzoeken = []
        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.srv.server_address[1]}"

    def stop(self):
        self.srv.shutdown()


class TestKern(unittest.TestCase):
    def test_sse(self):
        regels = ['data: {"t":"conv","id":"c1"}', "", 'data: {"t":"delta","text":"Hoi"}', "", 'data: {"t":"done","reply":"Hoi daar"}', ""]
        self.assertEqual([e["t"] for e in core.sse_events(regels)], ["conv", "delta", "done"])

    def test_config(self):
        pad = Path(tempfile.mkdtemp()) / "c.json"
        cfg = core.laad_config(pad)
        self.assertEqual(cfg["model"], "qwen3.5:9b")
        cfg["eco_sleutel"] = "x"
        core.bewaar_config(cfg, pad)
        self.assertEqual(core.laad_config(pad)["eco_sleutel"], "x")


class TestEcoEnBrug(unittest.TestCase):
    def setUp(self):
        self.opdrachten = [{"id": "11111111-2222-3333-4444-555555555555", "actie": "info", "args": {}}]
        self.klaar = []

        def behandel(h, methode, body):
            if h.headers.get("x-jarvis-key") != GEHEIM:
                return h.antwoord(401, {"error": "Geen toegang"})
            if isinstance(body, dict) and "nova" in body:
                nv = body["nova"]
                if nv["actie"] == "wacht":
                    o, self.opdrachten = self.opdrachten, []
                    return h.antwoord(200, {"opdrachten": o})
                if nv["actie"] == "klaar":
                    self.klaar.append(nv)
                    return h.antwoord(200, {"ok": True})
                if nv["actie"] == "geheugen":
                    return h.antwoord(200, {"feiten": ["Serkan studeert aan de HHS"], "agenda": "morgen Water 08:45", "lijstjes": ""})
            if isinstance(body, dict) and "tts" in body:
                return h.antwoord(200, b"\x00\x01" * 100, "audio/l16;rate=24000", {"x-audio-rate": "24000"})
            if isinstance(body, dict) and body.get("stream"):
                sse = "".join(f"data: {json.dumps(e)}\n\n" for e in ({"t": "conv", "id": "conv-1"}, {"t": "tool", "name": "pc"}, {"t": "delta", "text": "Spotify "}, {"t": "delta", "text": "staat aan."}, {"t": "done"}))
                return h.antwoord(200, sse.encode(), "text/event-stream")
            return h.antwoord(400, {"error": "?"})

        self.srv = NepServer(behandel)
        self.cfg = {**core.STANDAARD, "eco_url": self.srv.url, "eco_sleutel": GEHEIM}

    def tearDown(self):
        self.srv.stop()

    def test_chat_stroom(self):
        delen, tools = [], []
        antwoord, conv = core.Eco(self.cfg).chat("open spotify", on_delta=delen.append, on_tool=tools.append)
        self.assertEqual((antwoord, conv, tools), ("Spotify staat aan.", "conv-1", ["pc"]))
        self.assertEqual("".join(delen), "Spotify staat aan.")

    def test_tts_en_geheugen(self):
        eco = core.Eco(self.cfg)
        rate, r = eco.tts_stroom("hallo")
        with r:
            self.assertEqual((len(b"".join(r.iter_content(4800))), rate), (200, 24000))
        self.assertIn("HHS", eco.geheugen()["feiten"][0])

    def test_brug_voert_uit_en_meldt(self):
        gezien, online = [], []
        brug = core.Brug(core.Eco(self.cfg), lambda a, args: {"status": "gedaan", "actie": a}, lambda: {"naam": "test"},
                         on_online=lambda o, r: online.append(o), on_opdracht=gezien.append)
        self.assertEqual(brug.een_ronde(ms=0), 1)
        self.assertEqual(online, [True])
        self.assertEqual(self.klaar[0]["result"], {"status": "gedaan", "actie": "info"})
        self.assertEqual(gezien[0]["actie"], "info")

    def test_brug_kapotte_actie_stopt_niet(self):
        def stuk(a, args):
            raise RuntimeError("kapot")
        brug = core.Brug(core.Eco(self.cfg), stuk, lambda: {})
        brug.een_ronde(ms=0)
        self.assertEqual(self.klaar[0]["result"]["status"], "mislukt")

    def test_verkeerde_sleutel_is_offline(self):
        online = []
        brug = core.Brug(core.Eco({**self.cfg, "eco_sleutel": "fout"}), lambda a, b: {}, lambda: {}, on_online=lambda o, r: online.append((o, r)))
        brug.een_ronde(ms=0)
        self.assertFalse(online[0][0])
        self.assertIn("sleutel", online[0][1])

    def test_geen_server_is_offline(self):
        online = []
        brug = core.Brug(core.Eco({**self.cfg, "eco_url": "http://127.0.0.1:9"}), lambda a, b: {}, lambda: {}, on_online=lambda o, r: online.append(o))
        brug.een_ronde(ms=0)
        self.assertEqual(online, [False])


class TestLokaal(unittest.TestCase):
    def test_tool_lus_met_ollama(self):
        stap = {"n": 0}

        def behandel(h, methode, body):
            if h.path == "/api/tags":
                return h.antwoord(200, {"models": [{"name": "qwen3.5:9b"}]})
            if h.path == "/api/chat":
                stap["n"] += 1
                self.assertFalse(body["think"])
                if stap["n"] == 1:
                    return h.antwoord(200, {"message": {"role": "assistant", "content": "", "tool_calls": [{"function": {"name": "pc", "arguments": {"actie": "info"}}}]}})
                self.assertEqual(body["messages"][-1]["role"], "tool")
                return h.antwoord(200, {"message": {"role": "assistant", "content": "<think>even rekenen</think>Je batterij staat op 80 procent."}})
            return h.antwoord(404, {})

        srv = NepServer(behandel)
        try:
            gedaan = []
            lok = core.Lokaal({**core.STANDAARD, "ollama_url": srv.url}, lambda a, args: (gedaan.append(a), {"status": "gedaan", "batterij": "80%"})[1])
            lok.kennis = {"feiten": ["Serkan studeert aan de HHS"]}
            self.assertEqual(lok.status(), (True, "qwen3.5:9b"))
            self.assertEqual(lok.chat("hoe vol is mijn batterij?"), "Je batterij staat op 80 procent.")
            self.assertEqual(gedaan, ["info"])
            self.assertIn("HHS", lok.systeem())
        finally:
            srv.stop()

    def test_ollama_niet_actief(self):
        ok, reden = core.Lokaal({**core.STANDAARD, "ollama_url": "http://127.0.0.1:9"}, None).status()
        self.assertFalse(ok)
        self.assertIn("Ollama", reden)


class TestPc(unittest.TestCase):
    def test_url(self):
        self.assertEqual(pc.normaliseer_url("youtube"), "https://www.youtube.com")
        self.assertEqual(pc.normaliseer_url("nos.nl"), "https://nos.nl")
        self.assertTrue(pc.normaliseer_url("weer den haag").startswith("https://www.google.com/search?q=weer+den+haag"))

    def test_beste_match(self):
        k = [Path("Spotify.lnk"), Path("Microsoft Word.lnk"), Path("Discord.lnk")]
        self.assertEqual(pc.beste_match("word", k).stem, "Microsoft Word")
        self.assertIsNone(pc.beste_match("photoshop", k))

    def test_bestand_zoeken(self):
        d = Path(tempfile.mkdtemp())
        (d / "verslag water.docx").write_text("x")
        (d / "boodschappen.txt").write_text("x")
        self.assertEqual([p.name for p in pc.bestand_zoeken("verslag", [d])], ["verslag water.docx"])

    def test_onbekende_actie(self):
        self.assertEqual(pc.maak_uitvoerder()("rm -rf", {})["status"], "mislukt")


class TestApp(unittest.TestCase):
    """De logica van het venster (NovaApp) met een nep-venster en een nep-ECO."""

    def test_vraag_stroom(self):
        import nova

        class NepUI:
            def __init__(self):
                self.ev = []

            def __getattr__(self, naam):
                return lambda *a: self.ev.append((naam,) + a)

        ui = NepUI()
        app = nova.NovaApp(ui)
        app.cfg.update(eco_sleutel="x", stem=False)
        app.brug.online = True
        app.eco.chat = lambda t, c, on_delta, on_tool, on_ui=None, on_werk=None: (on_tool("pc"), on_delta("Spotify "), on_delta("staat aan."), ("Spotify staat aan.", "c1"))[-1]
        app.vraag("open spotify")
        self.assertEqual([e[0] for e in ui.ev], ["staat", "jij", "tool", "nova_start", "delta", "delta", "nova_eind", "staat"])
        self.assertEqual(app.conv, "c1")

    def test_spreken_eindigt_in_stand_by(self):
        ui, app = nep_app(online=True)
        app.eco.tts_stroom = lambda t: (24000, NepAntwoord([b"\x00\x00" * 10]))
        with nep_geluid():
            app.spreek("hallo")
            wacht_op(lambda: ui.ev and ui.ev[-1] == ("staat", "idle"))
        self.assertEqual([e for e in ui.ev if e[0] == "staat"], [("staat", "speaking"), ("staat", "idle")])

    def test_interface_bestanden(self):
        import nova
        ui = nova.ui_map()
        for f in ("index.html", "nova.js", "nova.css"):
            self.assertTrue((ui / f).exists(), f)
        html = (ui / "index.html").read_text(encoding="utf-8")
        self.assertIn("<b>E.C.O.", html)  # de laptop-app heet ECO; het offline brein N.O.V.A.
        self.assertIn('data-brein="nova"', html)


class NepUI:
    def __init__(self):
        self.ev = []

    def __getattr__(self, naam):
        return lambda *a: self.ev.append((naam,) + a)


def nep_app(online, map_=None):
    import nova
    core.MAP = map_ or Path(tempfile.mkdtemp())  # eigen map: nooit de echte instellingen of geschiedenis
    ui = NepUI()
    app = nova.NovaApp(ui)
    app.cfg.update(eco_sleutel="x", stem=True)
    app.brug.online = online
    return ui, app


def wacht_op(f, s=2.0):
    import time
    eind = time.time() + s
    while not f() and time.time() < eind:
        time.sleep(0.02)


class NepAntwoord:
    """Een gestreamd antwoord van ECO's stem; `fout` = de stroom breekt na de stukjes af."""

    def __init__(self, stukjes, fout=None):
        self.stukjes, self.fout, self.dicht = stukjes, fout, False

    def iter_content(self, n):
        yield from self.stukjes
        if self.fout:
            raise self.fout

    def __enter__(self):
        return self

    def __exit__(self, *a):
        self.dicht = True


class NepStroom:
    gespeeld = []

    def __init__(self, samplerate, channels, dtype):
        self.rate = samplerate

    def start(self): pass
    def stop(self): pass
    def abort(self): pass
    def close(self): pass

    def write(self, data):
        NepStroom.gespeeld.append(bytes(data))


def nep_geluid():
    NepStroom.gespeeld = []
    return mock.patch.dict(sys.modules, {"sounddevice": types.SimpleNamespace(RawOutputStream=NepStroom, stop=lambda: None)})


class NepMotor:
    def __init__(self, stemmen):
        self.stemmen, self.gezegd, self.instelling = stemmen, [], {}

    def getProperty(self, k): return self.stemmen
    def setProperty(self, k, v): self.instelling[k] = v
    def say(self, t): self.gezegd.append(t)
    def runAndWait(self): pass


def nep_pyttsx3(motor):
    return mock.patch.dict(sys.modules, {"pyttsx3": types.SimpleNamespace(init=lambda: motor)})


ENGELS = types.SimpleNamespace(id=r"HKEY_LOCAL_MACHINE\SOFTWARE\Microsoft\Speech\Voices\Tokens\TTS_MS_EN-US_ZIRA_11.0", name="Microsoft Zira Desktop - English (United States)", languages=[])
NEDERLANDS = types.SimpleNamespace(id=r"HKEY_LOCAL_MACHINE\SOFTWARE\Microsoft\Speech\Voices\Tokens\TTS_MS_NL-NL_FRANK_11.0", name="Microsoft Frank - Dutch (Netherlands)", languages=[])


class TestBrein(unittest.TestCase):
    """Verbonden met ECO of met NOVA: Serkan kiest; is het gekozen brein er niet, dan neemt het andere het over."""

    def setUp(self):
        self.pad = Path(tempfile.mkdtemp()) / "config.json"
        self.oud = core.bewaar_config
        core.bewaar_config = lambda cfg, pad=None: self.oud(cfg, self.pad)  # nooit de echte instellingen aanraken

    def tearDown(self):
        core.bewaar_config = self.oud

    def test_keuze_en_overnemen(self):
        ui, app = nep_app(online=True)
        app.lokaal_ok, app.lokaal_reden = True, "qwen3.5:9b"
        self.assertEqual(app.actief_brein(), ("eco", ""))
        app.kies_brein("nova")
        self.assertEqual(app.actief_brein(), ("nova", ""))
        status = [e[1] for e in ui.ev if e[0] == "status"][-1]
        self.assertEqual((status["titel"], status["gekozen"], status["modus"]), ("VERBONDEN MET N.O.V.A.", "nova", "online"))
        app.lokaal_ok, app.lokaal_reden = False, "Ollama draait niet"
        brein, reden = app.actief_brein()
        self.assertEqual(brein, "eco")
        self.assertIn("Ollama", reden)
        app.toon_status()
        status = [e[1] for e in ui.ev if e[0] == "status"][-1]
        self.assertEqual((status["titel"], status["modus"]), ("VERBONDEN MET ECO", "offline"))  # oranje: niet jouw keuze
        self.assertEqual(core.laad_config(self.pad)["brein"], "nova")

    def test_geen_brein(self):
        ui, app = nep_app(online=False)
        app.brug.reden = "geen internet"
        self.assertEqual(app.actief_brein()[0], "")
        app.toon_status()
        status = [e[1] for e in ui.ev if e[0] == "status"][-1]
        self.assertEqual((status["titel"], status["modus"]), ("GEEN BREIN BESCHIKBAAR", "geen"))
        self.assertNotIn("Offline-reserve", json.dumps(status))

    def test_nova_antwoordt_ook_online(self):
        ui, app = nep_app(online=True)
        app.cfg.update(stem=False, brein="nova")
        app.lokaal_ok = True
        app.eco.chat = lambda *a, **k: self.fail("ECO mag niet antwoorden als NOVA gekozen is")
        app.lokaal.chat = lambda t, on_tool=None: "Hier is NOVA."
        app.vraag("hoi")
        self.assertIn(("nova_eind", "Hier is NOVA."), ui.ev)

    def test_nova_luistert_met_whisper_van_eco(self):
        ui, app = nep_app(online=True)
        app.cfg.update(stem=False, brein="nova")
        app.lokaal_ok = True
        app.eco.spraak_chat = lambda *a, **k: self.fail("ECO mag niet antwoorden als NOVA gekozen is")
        app.eco.uitschrijven = lambda wav: "open spotify"
        app.lokaal.chat = lambda t, on_tool=None: f"NOVA: {t}"
        with mock.patch.object(stem, "opnemen", lambda: np.ones(16000, np.int16)), mock.patch.object(stem, "stop_spelen", lambda: None):
            app.luister()
        self.assertIn(("jij", "open spotify"), ui.ev)
        self.assertIn(("nova_eind", "NOVA: open spotify"), ui.ev)


class TestUpdateVragen(unittest.TestCase):
    """Zodra NOVA verbonden is en een nieuwe versie ziet, vraagt hij het; bij "ja" werkt hij meteen bij, bij "later" na 4 uur weer."""

    def setUp(self):
        import nova_update as upd
        self.ui, self.app = nep_app(online=True)
        self.app.cfg.update(stem=False)
        self.info = {"build": 99, "tag": "nova-99", "url": "x", "sha256": "", "grootte": 0}
        self.p = [mock.patch.object(upd, "nieuwste", lambda: self.info), mock.patch.object(upd, "huidige_build", lambda m: 12)]
        for x in self.p:
            x.start()

    def tearDown(self):
        for x in self.p:
            x.stop()

    def test_vraagt_een_keer(self):
        self.assertTrue(self.app.update_controleren())
        self.assertIn(("update", "nova-99", None), self.ui.ev)
        self.assertFalse(self.app.update_controleren())  # niet twee keer vragen voor dezelfde versie

    def test_ja_werkt_meteen_bij(self):
        self.app.update_controleren()
        gedaan = []
        with mock.patch.object(self.app, "update_uitvoeren", lambda: gedaan.append(True)):
            self.app.stuur("Ja, doe maar.")
            wacht_op(lambda: gedaan)
        self.assertEqual(gedaan, [True])
        self.assertIn(("jij", "Ja, doe maar."), self.ui.ev)

    def test_later_vraagt_pas_na_vier_uur(self):
        self.app.update_controleren()
        self.app.stuur("later")
        self.assertIsNone(self.app.update_info)
        self.assertIn(("update", None, None), self.ui.ev)
        self.assertFalse(self.app.update_controleren())
        self.app.update_later_tot = 0
        self.assertTrue(self.app.update_controleren())

    def test_gewone_vraag_gaat_gewoon_door(self):
        self.app.update_controleren()
        gevraagd = []
        with mock.patch.object(self.app, "vraag", lambda t: gevraagd.append(t)):
            self.app.stuur("wat is het weer")
            wacht_op(lambda: gevraagd)
        self.assertEqual(gevraagd, ["wat is het weer"])
        self.assertIsNotNone(self.app.update_info)  # de vraag over bijwerken blijft staan

    def test_geen_nieuwe_versie(self):
        self.info = {**self.info, "build": 12}
        self.assertFalse(self.app.update_controleren())
        self.assertNotIn("update", [e[0] for e in self.ui.ev])

    def test_bijwerken_meldt_voortgang(self):
        import hashlib
        import nova_update as upd
        data = b"N" * 300000

        def behandel(h, methode, body):
            return h.antwoord(200, data, "application/octet-stream")
        srv = NepServer(behandel)
        try:
            exe = Path(tempfile.mkdtemp()) / "NOVA.exe"
            exe.write_bytes(b"oud")
            stappen = []
            upd.bijwerken(exe, 1, {"build": 2, "url": srv.url + "/NOVA.exe", "sha256": hashlib.sha256(data).hexdigest(), "grootte": len(data)},
                          herstart=False, on_voortgang=stappen.append)
            self.assertEqual(exe.read_bytes(), data)
            self.assertEqual(stappen[-1], 1.0)
            self.assertGreater(len(stappen), 2)
        finally:
            srv.stop()


class TestInstellingen(unittest.TestCase):
    """Instellingen: geluid, beeld, AI-volgorde en indeling worden bewaard, en de volgordes worden echt gebruikt."""

    def setUp(self):
        self.pad = Path(tempfile.mkdtemp()) / "config.json"
        self.oud = core.bewaar_config
        core.bewaar_config = lambda cfg, pad=None: self.oud(cfg, self.pad)
        self.ui, self.app = nep_app(online=True)
        self.app.cfg.update(stem=False)

    def tearDown(self):
        core.bewaar_config = self.oud

    def test_volgorde_helper(self):
        self.assertEqual(core.volgorde({}, "stem"), ["gemini", "piper", "windows"])
        self.assertEqual(core.volgorde({"volgorde": {"stem": ["piper", "onzin", "piper"]}}, "stem"), ["piper", "gemini", "windows"])
        self.assertEqual(core.volgorde({"brein": "nova"}, "brein"), ["nova", "eco"])  # oude instelling telt nog

    def test_bewaren_en_controleren(self):
        self.app.instellingen({"volume": 250, "beeld": {"verberg": ["ollama", "pc"]}, "magneet": False, "geheim": "nee",
                               "volgorde": {"brein": ["nova", "eco"], "eco": ["gemini", "groq"], "nep": ["x"]},
                               "indeling": {"kern": {"x": 0.3, "y": 0.1, "w": 0.4, "h": 0.6, "z": 3}}})
        cfg = core.laad_config(self.pad)
        self.assertEqual(cfg["volume"], 100)  # begrensd
        self.assertEqual(cfg["beeld"]["verberg"], ["ollama", "pc"])
        self.assertFalse(cfg["magneet"])
        self.assertNotIn("geheim", cfg)
        self.assertNotIn("nep", cfg["volgorde"])
        self.assertEqual((cfg["brein"], self.app.gekozen()), ("nova", "nova"))
        self.assertEqual(cfg["indeling"]["kern"]["w"], 0.4)
        inst = self.app.inst()
        self.assertEqual(inst["volgorde"]["eco"], ["gemini", "groq", "claude"])
        self.assertEqual(self.app.info()["inst"]["volume"], 100)

    def test_eco_krijgt_de_modelvolgorde(self):
        self.app.instellingen({"volgorde": {"eco": ["claude", "gemini"]}})
        self.assertEqual(self.app.eco._ui()["modellen"], ["claude", "gemini", "groq"])

    def test_stem_volgt_de_volgorde(self):
        self.app.instellingen({"volgorde": {"stem": ["piper", "gemini", "windows"]}})
        self.app.eco.tts_stroom = lambda t: self.fail("Gemini mag niet als eerste als Piper bovenaan staat")
        gezegd = []
        with mock.patch.object(stem, "piper_klaar", lambda: True), mock.patch.object(stem, "piper_spreek", gezegd.append):
            self.assertEqual(self.app.spreek_met_stem("hoi"), ("PIPER (NATHALIE)", ""))
        self.assertEqual(gezegd, ["hoi"])

    def test_volume_maakt_zachter(self):
        self.app.instellingen({"volume": 50})
        self.app.eco.tts_stroom = lambda t: (24000, NepAntwoord([np.full(24000, 1000, np.int16).tobytes()]))
        with nep_geluid():
            self.app.spreek_met_stem("hoi")
        gespeeld = np.frombuffer(b"".join(NepStroom.gespeeld), np.int16)
        self.assertEqual(int(gespeeld.max()), 500)
        stem.VOLUME = 1.0

    def test_spraak_volgorde_vosk_eerst(self):
        self.app.instellingen({"volgorde": {"spraak": ["vosk", "whisper"]}})
        self.app.eco.spraak_chat = lambda *a, **k: self.fail("met Vosk bovenaan niet via ECO's spraak")
        self.app.eco.uitschrijven = lambda wav: self.fail("Vosk staat bovenaan")
        gevraagd = []
        with mock.patch.object(stem, "opnemen", lambda: np.ones(16000, np.int16)), mock.patch.object(stem, "stop_spelen", lambda: None), \
                mock.patch.object(stem, "vosk_klaar", lambda: True), mock.patch.object(stem, "vosk_herken", lambda s: "hoe laat is het"), \
                mock.patch.object(self.app, "vraag", lambda t, al_getoond=False: gevraagd.append(t)):
            self.app.luister()
        self.assertEqual(gevraagd, ["hoe laat is het"])


class TestKlappen(unittest.TestCase):
    """2× klappen = ECO luistert. Nagebootst geluid: klappen, praten, muziek, applaus."""
    R = 16000

    def setUp(self):
        self.rng = np.random.default_rng(1)

    def stil(self, s, n=0.002):
        return self.rng.normal(0, n, int(self.R * s))

    def klap(self, luid=0.5):  # 2 ms aanzet, ~25 ms uitsterven, alle tonen door elkaar
        t = np.arange(int(self.R * 0.08)) / self.R
        return self.rng.normal(0, 1, len(t)) * luid * np.minimum(1, t / 0.002) * np.exp(-t / 0.025)

    def stem(self, s):
        t = np.arange(int(self.R * s)) / self.R
        return 0.15 * (np.sin(2 * np.pi * 180 * t) + 0.5 * np.sin(2 * np.pi * 420 * t)) * (0.6 + 0.4 * np.sin(2 * np.pi * 4 * t))

    def muziek(self, s):  # beat (tik) elke 0,5 s over een toon
        t = np.arange(int(self.R * s)) / self.R
        x = 0.1 * np.sin(2 * np.pi * 220 * t)
        for b in np.arange(0, s, 0.5):
            i, k = int(b * self.R), self.klap(0.4)
            x[i:i + len(k)] += k[:len(x) - i]
        return x

    def pcm(self, delen):
        return (np.clip(np.concatenate(delen), -1, 1) * 32767).astype(np.int16)

    def treffers(self, delen):
        import nova_klap
        d, x = nova_klap.KlapDetector(self.R), self.pcm(delen)
        return sum(d.voed(x[i:i + 512]) for i in range(0, len(x), 512))

    def test_twee_klappen(self):
        self.assertEqual(self.treffers([self.stil(1), self.klap(), self.stil(0.3), self.klap(), self.stil(1)]), 1)
        self.assertEqual(self.treffers([self.stil(1, 0.004), self.klap(0.15), self.stil(0.4, 0.004), self.klap(0.15), self.stil(1, 0.004)]), 1)
        self.assertEqual(self.treffers([self.stem(2), self.stil(0.6), self.klap(), self.stil(0.35), self.klap(), self.stil(1)]), 1)

    def test_geen_klappen(self):
        for naam, delen in {"één klap": [self.stil(1), self.klap(), self.stil(2)],
                            "drie klappen": [self.stil(1), self.klap(), self.stil(0.25), self.klap(), self.stil(0.25), self.klap(), self.stil(1)],
                            "te ver uit elkaar": [self.stil(1), self.klap(), self.stil(1.5), self.klap(), self.stil(1)],
                            "echo": [self.stil(1), self.klap(), self.stil(0.03), self.klap(), self.stil(1)],
                            "praten": [self.stil(0.5), self.stem(3), self.stil(1)],
                            "muziek": [self.muziek(6)]}.items():
            with self.subTest(naam):
                self.assertEqual(self.treffers(delen), 0)

    def nep_microfoon(self, x):
        class Microfoon:
            def __init__(self, samplerate, channels, dtype, blocksize, callback):
                self.cb, self.n = callback, blocksize

            def __enter__(self):
                def speel():
                    for i in range(0, len(x), self.n):
                        self.cb(x[i:i + self.n].reshape(-1, 1), self.n, None, None)
                threading.Thread(target=speel, daemon=True).start()
                return self

            def __exit__(self, *a):
                pass
        return mock.patch.dict(sys.modules, {"sounddevice": types.SimpleNamespace(InputStream=Microfoon)})

    def test_klappen_start_luisteren(self):
        ui, app = nep_app(online=True)
        geluisterd = []
        with self.nep_microfoon(self.pcm([self.stil(1), self.klap(), self.stil(0.3), self.klap(), self.stil(1)])), \
                mock.patch.object(app, "luister", lambda: geluisterd.append(1)):
            app.klap_lus(rondes=1)
        self.assertEqual(geluisterd, [1])
        self.assertIn(("sys", "👏👏 Ik luister."), ui.ev)

    def test_niet_tijdens_antwoord_of_praten(self):
        _, app = nep_app(online=True)
        self.assertTrue(app.klap_mag())
        app.bezig = True
        self.assertFalse(app.klap_mag())
        app.bezig = False
        with app.spreek_slot:  # ECO praat: dan hoort hij zichzelf niet
            self.assertFalse(app.klap_mag())
        app.spreekt = True  # het antwoord is klaar en de stem start zo: ook dan niet
        self.assertFalse(app.klap_mag())
        app.spreekt = False
        app.cfg["klappen"] = False
        self.assertFalse(app.klap_mag())

    def test_geen_microfoon(self):
        ui, app = nep_app(online=True)

        def kapot(*a, **k):
            raise OSError("geen invoerapparaat")
        with mock.patch.dict(sys.modules, {"sounddevice": types.SimpleNamespace(InputStream=kapot)}), mock.patch("time.sleep", lambda s: None):
            app.klap_lus(rondes=1)
        self.assertIn("geen invoerapparaat", app.klap_fout)
        app.toon_status()
        self.assertEqual([e for e in ui.ev if e[0] == "status"][-1][1]["klappen"], "GEEN MICROFOON")

    def test_eco_zet_klappen_uit(self):  # "zet klappen uit" via ECO's tool interface
        pad = Path(tempfile.mkdtemp()) / "config.json"
        oud = core.bewaar_config
        core.bewaar_config = lambda cfg, p=None: oud(cfg, pad)
        try:
            ui, app = nep_app(online=True)
            app.cfg["stem"] = False
            regels = ['data: {"t":"tool","name":"interface"}', "", 'data: {"t":"ui","onderdeel":"klappen","aan":false}', "",
                      'data: {"t":"done","reply":"Klappen staat uit."}', ""]
            r = types.SimpleNamespace(iter_lines=lambda decode_unicode=True: iter(regels), encoding=None)
            with mock.patch.object(app.eco, "_post", lambda *a, **k: r):
                app.vraag("zet klappen uit")
            self.assertFalse(app.cfg["klappen"])
            self.assertFalse([e for e in ui.ev if e[0] == "inst"][-1][1]["klappen"])
            self.assertTrue(app.eco._ui()["klappen"] is False)
        finally:
            core.bewaar_config = oud

    def test_instelling(self):
        pad = Path(tempfile.mkdtemp()) / "config.json"
        oud = core.bewaar_config
        core.bewaar_config = lambda cfg, p=None: oud(cfg, pad)
        try:
            ui, app = nep_app(online=True)
            self.assertTrue(app.inst()["klappen"])  # standaard aan
            app.instellingen({"klappen": False})
            self.assertFalse(core.laad_config(pad)["klappen"])
            self.assertFalse(app.inst()["klappen"])
            self.assertEqual([e for e in ui.ev if e[0] == "status"][-1][1]["klappen"], "UIT")
        finally:
            core.bewaar_config = oud


class TestWerkmodus(unittest.TestCase):
    """Werkmodus: live laten zien wat ECO doet, met een schermplaatje na een laptop-actie."""

    def test_pc_actie_met_beeld(self):
        ui, app = nep_app(online=True)
        app.uitvoeren = lambda actie, args: {"status": "gelukt", "controle": "Spotify staat open"}
        with mock.patch.object(pc, "schermpje_jpeg", lambda breed=560: "QUJD"), mock.patch("time.sleep", lambda s: None):
            app.pc_actie("app_openen", {"naam": "Spotify"})
            wacht_op(lambda: any(e[0] == "werk" and e[1]["t"] == "tool_beeld" for e in ui.ev))
        werk = [e[1] for e in ui.ev if e[0] == "werk"]
        self.assertEqual([w["t"] for w in werk], ["tool", "tool_klaar", "tool_beeld"])
        self.assertEqual(werk[0]["info"], "app openen Spotify")
        self.assertTrue(werk[1]["ok"])
        self.assertEqual(werk[1]["uit"], "Spotify staat open")
        self.assertEqual(werk[2]["beeld"], "QUJD")
        self.assertEqual(len({w["id"] for w in werk}), 1)

    def test_mislukt_en_uit(self):
        ui, app = nep_app(online=True)
        app.uitvoeren = lambda actie, args: {"status": "mislukt", "fout": "app niet gevonden"}
        app.pc_actie("vergrendelen", {})
        klaar = [e[1] for e in ui.ev if e[0] == "werk" and e[1]["t"] == "tool_klaar"][0]
        self.assertFalse(klaar["ok"])
        self.assertEqual(klaar["uit"], "app niet gevonden")
        ui.ev.clear()
        app.cfg["werkmodus"] = False
        app.pc_actie("vergrendelen", {})
        self.assertFalse([e for e in ui.ev if e[0] == "werk"])
        self.assertTrue([e for e in ui.ev if e[0] == "pc"])  # het PC-paneel blijft gewoon werken

    def test_stappen_van_eco(self):
        ui, app = nep_app(online=True)
        app.cfg["stem"] = False
        regels = ['data: {"t":"tool","name":"web_zoeken","id":"w1","info":"weer Den Haag"}', "",
                  'data: {"t":"tool_klaar","name":"web_zoeken","id":"w1","ok":true,"uit":"12 graden","bronnen":["KNMI"]}', "",
                  'data: {"t":"tool","name":"pc","id":"w2","info":"app_openen"}', "",
                  'data: {"t":"done","reply":"Het is 12 graden."}', ""]
        r = types.SimpleNamespace(iter_lines=lambda decode_unicode=True: iter(regels), encoding=None)
        with mock.patch.object(app.eco, "_post", lambda *a, **k: r):
            app.vraag("wat voor weer is het")
        werk = [e[1] for e in ui.ev if e[0] == "werk"]
        self.assertEqual([(w["t"], w["name"]) for w in werk], [("tool", "web_zoeken"), ("tool_klaar", "web_zoeken")])  # pc toont de laptop zelf
        self.assertEqual(werk[1]["bronnen"], ["KNMI"])

    def test_instelling(self):
        pad = Path(tempfile.mkdtemp()) / "config.json"
        oud = core.bewaar_config
        core.bewaar_config = lambda cfg, p=None: oud(cfg, pad)
        try:
            _, app = nep_app(online=True)
            self.assertTrue(app.inst()["werkmodus"])
            app.instellingen({"werkmodus": False})
            self.assertFalse(core.laad_config(pad)["werkmodus"])
        finally:
            core.bewaar_config = oud


class TestOnthouden(unittest.TestCase):
    """NOVA onthoudt het transcript, de PC-acties, het gesprek en de plek van het venster als hij uit en aan gaat."""

    def test_transcript_en_gesprek_na_herstart(self):
        map_ = Path(tempfile.mkdtemp())
        ui, app = nep_app(online=True, map_=map_)
        app.cfg.update(stem=False)
        app.eco.chat = lambda t, c, on_delta, on_tool, on_ui=None, on_werk=None: (on_delta("Het is 13:05."), ("Het is 13:05.", "conv-7"))[-1]
        app.vraag("hoe laat is het")
        app.pc_actie("info", {})
        # NOVA uit en weer aan
        ui2, app2 = nep_app(online=True, map_=map_)
        self.assertEqual(app2.conv, "conv-7")  # het gesprek met ECO gaat verder
        app2.start()
        app2.stop()
        herstel = next(e[1] for e in ui2.ev if e[0] == "herstel")
        self.assertEqual([(r["rol"], r["tekst"]) for r in herstel["transcript"]], [("user", "hoe laat is het"), ("assistant", "Het is 13:05.")])
        self.assertEqual(herstel["pc"][0]["tekst"], "info")

    def test_oud_gesprek_begint_opnieuw(self):
        import nova
        log = nova.Logboek(Path(tempfile.mkdtemp()) / "g.json")
        log.zet_gesprek("conv-1", [{"role": "user", "content": "hoi"}])
        self.assertEqual(nova.Logboek(log.pad).gesprek(), "conv-1")
        log.data["conv_tijd"] = 0
        log.bewaar()
        self.assertIsNone(nova.Logboek(log.pad).gesprek())
        self.assertEqual(nova.Logboek(log.pad).data["lokaal"][0]["content"], "hoi")

    def test_maximum_en_wissen(self):
        import nova
        log = nova.Logboek(Path(tempfile.mkdtemp()) / "g.json")
        for i in range(250):
            log.regel("user", f"vraag {i}")
        self.assertEqual(len(nova.Logboek(log.pad).data["transcript"]), 200)
        log.wis()
        self.assertEqual(nova.Logboek(log.pad).data["transcript"], [])

    def test_kapot_bestand(self):
        import nova
        pad = Path(tempfile.mkdtemp()) / "g.json"
        pad.write_text("{kapot", encoding="utf-8")
        self.assertEqual(nova.Logboek(pad).data["transcript"], [])

    def test_venster_plek(self):
        import nova
        sc = [types.SimpleNamespace(x=0, y=0, width=1920, height=1080)]
        self.assertEqual(nova.venster_geometrie({"x": 100, "y": 50, "w": 1500, "h": 900}, sc), {"width": 1500, "height": 900, "maximized": False, "x": 100, "y": 50})
        self.assertNotIn("x", nova.venster_geometrie({"x": 3000, "y": 50, "w": 1500, "h": 900}, sc))  # tweede scherm is weg: midden in beeld
        self.assertTrue(nova.venster_geometrie({"max": True}, sc)["maximized"])
        self.assertEqual(nova.venster_geometrie({"x": 0, "y": 0, "w": 5000, "h": 300}, sc)["width"], 1920)

    def test_venster_bewaren(self):
        ui, app = nep_app(online=False)
        app.venster(x=100, y=80)
        app.venster(w=1400, h=900)
        app.venster(x=-32000, y=-32000)  # geminimaliseerd: telt niet
        app.venster(max=True)
        app.venster(w=1920, h=1040)  # gemaximaliseerd: de gewone grootte blijft
        app.venster_bewaren()
        v = core.laad_config()["venster"]
        self.assertEqual((v["x"], v["y"], v["w"], v["h"], v["max"]), (100, 80, 1400, 900, True))


class TestNovaStem(unittest.TestCase):
    """Welke stem NOVA gebruikt: ECO's stem, dan Piper, dan een Nederlandse Windows-stem; nooit een Engelse."""

    def setUp(self):
        stem.nieuw_spreken()

    def test_eco_stem_gestreamd(self):
        ui, app = nep_app(online=True)
        stukjes = [b"\x01" * 4801, b"\x02" * 9599, b"\x03" * 7]  # rare grenzen, midden in samples
        app.eco.tts_stroom = lambda t: (24000, NepAntwoord(stukjes))
        with nep_geluid():
            self.assertEqual(app.spreek_met_stem("hallo"), ("GEMINI (ECO)", ""))
        gespeeld = b"".join(NepStroom.gespeeld)
        self.assertEqual(gespeeld, b"".join(stukjes)[:len(gespeeld)])
        self.assertEqual(len(gespeeld), 14406)  # alles, op het halve sample na
        self.assertTrue(all(len(d) % 2 == 0 for d in NepStroom.gespeeld))

    def test_eco_stem_mislukt_dan_piper(self):
        ui, app = nep_app(online=True)

        def stuk(t):
            raise core.EcoFout("ECO 502: Gemini-stem niet beschikbaar")
        app.eco.tts_stroom = stuk
        gezegd = []
        with mock.patch.object(stem, "piper_klaar", lambda: True), mock.patch.object(stem, "piper_spreek", gezegd.append):
            naam, reden = app.spreek_met_stem("hallo")
            app.stem_gebruikt(naam, reden)
        self.assertEqual((naam, gezegd), ("PIPER (NATHALIE)", ["hallo"]))
        self.assertIn("Gemini-stem niet beschikbaar", reden)
        self.assertEqual(app.stem_status()["stem_soort"], "warn")
        self.assertTrue(any(e[0] == "sys" and "piper" in e[1] for e in ui.ev))

    def test_halverwege_weg_geen_tweede_stem(self):
        ui, app = nep_app(online=True)
        app.eco.tts_stroom = lambda t: (24000, NepAntwoord([b"\x00" * 48000], fout=ConnectionError("stroom brak af")))
        gezegd = []
        with nep_geluid(), mock.patch.object(stem, "piper_klaar", lambda: True), mock.patch.object(stem, "piper_spreek", gezegd.append):
            naam, reden = app.spreek_met_stem("hallo")
        self.assertEqual((naam, gezegd), ("GEMINI (ECO)", []))
        self.assertIn("halverwege", reden)

    def test_nooit_een_engelse_stem(self):
        ui, app = nep_app(online=False)
        motor = NepMotor([ENGELS])
        with nep_pyttsx3(motor), mock.patch.object(stem, "piper_klaar", lambda: False):
            naam, reden = app.spreek_met_stem("Goedemorgen Serkan")
            app.stem_gebruikt(naam, reden)
        self.assertEqual(motor.gezegd, [])
        self.assertEqual(naam, "GEEN — ALLEEN TEKST")
        self.assertIn("gedownload", reden)
        self.assertEqual(app.stem_status()["stem_soort"], "bad")

    def test_nederlandse_windows_stem(self):
        motor = NepMotor([ENGELS, NEDERLANDS])
        with nep_pyttsx3(motor):
            self.assertTrue(stem.windows_stem("hallo"))
        self.assertEqual((motor.gezegd, motor.instelling["voice"]), (["hallo"], NEDERLANDS.id))

    def test_onderbroken_geen_andere_stem(self):
        ui, app = nep_app(online=True)

        def onderbreek(t):
            stem.stop_spelen()
            return 24000, NepAntwoord([b"\x00" * 48000])
        app.eco.tts_stroom = onderbreek
        gezegd = []
        with nep_geluid(), mock.patch.object(stem, "piper_klaar", lambda: True), mock.patch.object(stem, "piper_spreek", gezegd.append):
            app.spreek_met_stem("hallo")
        self.assertEqual((NepStroom.gespeeld, gezegd), ([], []))

    def test_piper_downloaden(self):
        bestanden = {"/nl_BE-nathalie-medium.onnx": b"MODEL" * 1000, "/nl_BE-nathalie-medium.onnx.json": b'{"audio": {"sample_rate": 22050}}'}

        def behandel(h, methode, body):
            return h.antwoord(200, bestanden[h.path], "application/octet-stream") if h.path in bestanden else h.antwoord(404, {})
        srv = NepServer(behandel)
        try:
            with mock.patch.object(stem, "PIPER_MAP", Path(tempfile.mkdtemp()) / "stem"):
                self.assertFalse(stem.piper_klaar())
                voortgang = []
                stem.piper_downloaden(voortgang.append, url=srv.url + "/")
                self.assertTrue(stem.piper_klaar())
                self.assertEqual((stem.PIPER_MAP / "nl_BE-nathalie-medium.onnx").read_bytes(), bestanden["/nl_BE-nathalie-medium.onnx"])
                self.assertEqual(list(stem.PIPER_MAP.glob("*.deel")), [])
                self.assertEqual(voortgang[-1], 1.0)
        finally:
            srv.stop()

    @unittest.skipUnless(__import__("importlib").util.find_spec("piper"), "piper-tts niet geïnstalleerd")
    def test_piper_uitspraak(self):
        with mock.patch.object(stem, "PIPER_MAP", Path(tempfile.mkdtemp())):
            self.assertIn("uitspraak werkt", stem.piper_proef())


class TestDocumenten(unittest.TestCase):
    """Bestanden lezen (pdf, Word, Excel, PowerPoint, tekst) en een PDF of Word-document maken."""

    def setUp(self):
        import nova_doc
        self.doc = nova_doc
        self.map = Path(tempfile.mkdtemp())

    def test_maken_en_teruglezen(self):
        for soort, ext in (("pdf", ".pdf"), ("word", ".docx")):
            r = self.doc.document_maken("Samenvatting: water?", "# Kop\n- punt één\nGewone zin met €5.", soort, map_=self.map, openen=False)
            self.assertEqual((r["status"], Path(r["bestand"]).suffix), ("gedaan", ext))
            self.assertNotIn(":", Path(r["bestand"]).name)  # geen tekens die Windows niet wil
            tekst, _ = self.doc.lees(Path(r["bestand"]))
            self.assertIn("punt", tekst)
            self.assertIn("Gewone zin", tekst)

    def test_nooit_overschrijven(self):
        a = self.doc.document_maken("Brief", "eerste", "word", map_=self.map, openen=False)["bestand"]
        b = self.doc.document_maken("Brief", "tweede", "word", map_=self.map, openen=False)["bestand"]
        self.assertNotEqual(a, b)
        self.assertIn("eerste", self.doc.lees(Path(a))[0])

    def test_leeg_document(self):
        self.assertEqual(self.doc.document_maken("x", "  ", map_=self.map, openen=False)["status"], "mislukt")

    def test_excel_powerpoint_tekst_lezen(self):
        from openpyxl import Workbook
        from pptx import Presentation
        wb = Workbook()
        wb.active.append(["vak", "cijfer"])
        wb.active.append(["Water", 7.5])
        wb.save(self.map / "cijfers.xlsx")
        prs = Presentation()
        dia = prs.slides.add_slide(prs.slide_layouts[1])
        dia.shapes.title.text = "Presentatie water"
        prs.save(self.map / "presentatie.pptx")
        (self.map / "notities.txt").write_text("boodschappen: melk", encoding="utf-8")
        self.assertIn("Water | 7.5", self.doc.lees(self.map / "cijfers.xlsx")[0])
        self.assertIn("Presentatie water", self.doc.lees(self.map / "presentatie.pptx")[0])
        self.assertIn("melk", self.doc.lees(self.map / "notities.txt")[0])

    def test_bestand_lezen_via_pc(self):
        self.doc.document_maken("Verslag water", "# Inleiding\nHet water is schoon.", "pdf", map_=self.map, openen=False)
        (self.map / "verslag water.exe").write_bytes(b"MZ")
        with mock.patch.object(pc, "zoekmappen", lambda: [self.map]):
            r = pc.maak_uitvoerder()("bestand_lezen", {"naam": "verslag water"})
        self.assertEqual(r["status"], "gedaan")
        self.assertTrue(r["bestand"].endswith(".pdf"))  # het .exe-bestand slaat hij over
        self.assertIn("schoon", r["tekst"])
        self.assertIn("geen opdracht", r["let_op"])

    def test_niet_gevonden_of_gescand(self):
        from pypdf import PdfWriter
        w = PdfWriter()
        w.add_blank_page(200, 200)
        with open(self.map / "scan.pdf", "wb") as f:
            w.write(f)
        zoek = lambda n: [p for p in self.map.iterdir() if n in p.name]
        self.assertEqual(self.doc.bestand_lezen("bestaatniet", zoek)["status"], "mislukt")
        r = self.doc.bestand_lezen("scan", zoek)
        self.assertEqual(r["status"], "mislukt")
        self.assertIn("gescand", r["fout"])

    def test_document_maken_via_pc(self):
        with mock.patch.object(self.doc, "documentenmap", lambda: self.map), mock.patch.object(self.doc.os, "startfile", create=True):
            r = pc.maak_uitvoerder()("document_maken", {"naam": "Planning toetsweek", "tekst": "- maandag: water", "soort": "word"})
        self.assertEqual(r["status"], "gedaan")
        self.assertTrue(Path(r["bestand"]).is_file())

    def test_offline_brein_kapt_grote_documenten_af(self):
        stap = {"n": 0}

        def behandel(h, methode, body):
            if h.path == "/api/chat":
                stap["n"] += 1
                if stap["n"] == 1:
                    return h.antwoord(200, {"message": {"role": "assistant", "content": "", "tool_calls": [{"function": {"name": "pc", "arguments": {"actie": "bestand_lezen", "naam": "boek"}}}]}})
                self.assertLess(len(body["messages"][-1]["content"]), 9500)
                self.assertEqual(body["options"]["num_ctx"], 8192)
                return h.antwoord(200, {"message": {"role": "assistant", "content": "Het boek gaat over water."}})
            return h.antwoord(404, {})
        srv = NepServer(behandel)
        try:
            lok = core.Lokaal({**core.STANDAARD, "ollama_url": srv.url}, lambda a, args: {"status": "gedaan", "tekst": "x" * 15000})
            self.assertEqual(lok.chat("waar gaat mijn boek over?"), "Het boek gaat over water.")
        finally:
            srv.stop()


class TestUpdate(unittest.TestCase):
    def test_bijwerken(self):
        import hashlib
        import nova_update as upd
        nieuw_exe = b"NIEUWE NOVA" * 100

        def behandel(h, methode, body):
            if h.path == "/latest":
                return h.antwoord(200, {"tag_name": "nova-12", "assets": [{"name": "NOVA.exe", "browser_download_url": srv.url + "/NOVA.exe",
                                                                            "size": len(nieuw_exe), "digest": "sha256:" + hashlib.sha256(nieuw_exe).hexdigest()}]})
            if h.path == "/NOVA.exe":
                return h.antwoord(200, nieuw_exe, "application/octet-stream")
            return h.antwoord(404, {})

        srv = NepServer(behandel)
        try:
            d = Path(tempfile.mkdtemp())
            exe = d / "NOVA.exe"
            exe.write_bytes(b"OUDE NOVA")
            info = upd.nieuwste(srv.url + "/latest")
            self.assertEqual(info["build"], 12)
            self.assertFalse(upd.bijwerken(exe, 12, info, herstart=False))  # al de nieuwste
            self.assertTrue(upd.bijwerken(exe, 3, info, herstart=False))
            self.assertEqual(exe.read_bytes(), nieuw_exe)
            self.assertEqual((d / "NOVA.oud.exe").read_bytes(), b"OUDE NOVA")
            upd.opruimen(exe)
            self.assertFalse((d / "NOVA.oud.exe").exists())
            # verkeerde controlesom: niets vervangen
            info["sha256"] = "0" * 64
            exe.write_bytes(b"HUIDIG")
            with self.assertRaises(RuntimeError):
                upd.bijwerken(exe, 3, info, herstart=False)
            self.assertEqual(exe.read_bytes(), b"HUIDIG")
        finally:
            srv.stop()


class TestStem(unittest.TestCase):
    def blokken(self, *delen):
        for niveau, sec in delen:
            for _ in range(int(sec * 10)):
                yield (np.random.randn(1600) * niveau * 32767).astype(np.int16)

    def test_stopt_na_stilte(self):
        samples, gesproken = stem.stilte_detectie(self.blokken((0.002, 0.5), (0.2, 1.0), (0.002, 3.0)))
        self.assertTrue(gesproken)
        self.assertLess(len(samples) / stem.RATE, 3.0)  # gestopt na ~1,2 s stilte, niet na 4,5 s

    def test_niets_gezegd(self):
        _, gesproken = stem.stilte_detectie(self.blokken((0.002, 9.0)))
        self.assertFalse(gesproken)

    def test_wav(self):
        self.assertEqual(stem.naar_wav(np.zeros(160, np.int16))[:4], b"RIFF")


if __name__ == "__main__":
    unittest.main()
